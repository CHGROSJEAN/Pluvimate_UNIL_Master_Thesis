# Charlotte Grosjean - 25.03.2026 - UNIL Master Thesis
"""
CONSOLIDATE PROCESSED STATION DATA

This script:
- Reads processed station files (.txt, .xlsx, .csv)
- Cleans and standardizes precipitation data
- Fills missing timesteps (3 min resolution)
- Exports consolidated datasets
- Generates a global summary Excel per station
"""

# %% 0 - Imports and configuration

import os
import pandas as pd
from glob import glob
from A_data_processing_utils import parse_date, extract_numeric
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from I_paths_config import PROCESSED_DIRECTORY, CONSOLIDATED_DIRECTORY, AVAILABLE_DATA_FILE

base_directory_processed = PROCESSED_DIRECTORY
output_directory = CONSOLIDATED_DIRECTORY
os.makedirs(output_directory, exist_ok=True)


# %% 1 - Helper functions

def clean_and_prepare(df, file_name):
    """
    Clean a raw dataframe:
    - Extract numeric precipitation
    - Remove invalid rows (NaNs)
    """
    df['Precipitation'] = df['Precipitation'].apply(extract_numeric)
    valid = df.dropna(subset=['DateTime', 'Precipitation']).copy()
    valid['Source_File'] = file_name
    return valid[['Source_File', 'DateTime', 'Precipitation']]


def fill_missing_timesteps(df, expected_timestep_minutes=3):
    """
    Fill temporal gaps in a time series.
    -1: inter-file gap
    -2: intra-file gap
    """

    # Ensure DateTime is sorted
    df = df.sort_values('DateTime').reset_index(drop=True)
    expected_delta = pd.Timedelta(minutes=expected_timestep_minutes)
    all_rows = []
    prev_file = None
    last_time = None

    for _, row in df.iterrows():
        file_name = row['Source_File']
        dt = pd.to_datetime(row['DateTime'])
        precip = row['Precipitation']

        if last_time is not None:                                                               # Check for gaps
            gap = dt - last_time                                                                # If gap is larger than timesteps, fill in missing timesteps  
            if gap > expected_delta:
                missing_times = pd.date_range(
                    start=last_time + expected_delta,
                    end=dt - expected_delta,
                    freq=f'{expected_timestep_minutes}min'
                )
                for missing_dt in missing_times:                                                # complete the gap with data
                    gap_type = -1 if prev_file != file_name else -2                             # inter-file (-1) vs intra-file gap (-2)
                    all_rows.append({
                        'Source_File': "No data" if gap_type == -1 else file_name,
                        'DateTime': missing_dt,
                        'Precipitation': gap_type
                    })

        all_rows.append({'Source_File': file_name, 'DateTime': dt, 'Precipitation': precip})
        prev_file = file_name
        last_time = dt

    return pd.DataFrame(all_rows).sort_values('DateTime').reset_index(drop=True)


# %% 2 - Main function per station

def process_consolidate(directory):
    """Process and consolidate all files for a station.
    1. Read all files → clean and standardize
    2. Merge into one complete dataframe
    3. Fill missing timesteps
    4. Save consolidated CSV
    5. Compute summary statistics"""


    station_name = os.path.basename(directory)
    print(f"\nProcessing station: {station_name}")

    txt_files = glob(os.path.join(directory, "*.txt"))
    xlsx_files = glob(os.path.join(directory, "*.xlsx"))
    csv_files = glob(os.path.join(directory, "*.csv"))

    dataframes = []

    # --- TXT ---
    for file_path in txt_files:
        try:
            df = pd.read_csv(file_path, sep='\t', skiprows=5, header=None,              # 5 first row in .txt are metadata
                             usecols=[0, 1, 2],
                             names=['Index', 'DateTime', 'Precipitation'])
            df['DateTime'] = df['DateTime'].apply(parse_date)                           # parse multiple date formats
            valid = clean_and_prepare(df, os.path.basename(file_path))                  # function from previous cell to clean and prepare the dataframe
            dataframes.append(valid)
        except Exception as e:
            print(f"Error TXT {file_path}: {e}")

    # --- XLSX ---
    for file_path in xlsx_files:
        if os.path.basename(file_path).startswith('~$'):
            continue
        try:
            df = pd.read_excel(file_path, skiprows=6)                                   # 6 first row in .xlsx are metadata, the 7th row is the header    
            df.columns = df.columns.str.strip()
            date_col = next((c for c in df.columns if 'date' in c.lower() or 'time' in c.lower()), df.columns[0])
            precip_col = next((c for c in df.columns if 'precip' in c.lower() or 'pulse' in c.lower()), None)
            if precip_col:
                df['DateTime'] = df[date_col].apply(parse_date)
                df['DateTime'] = pd.to_datetime(df['DateTime'], errors='coerce')
                df['Precipitation'] = df[precip_col]
                valid = clean_and_prepare(df, os.path.basename(file_path))
                dataframes.append(valid)
        except Exception as e:
            print(f"Error XLSX {file_path}: {e}")

    # --- CSV ---
    for file_path in csv_files:
        try:
            df = pd.read_csv(file_path)
            if 'created_at' in df.columns and 'field5' in df.columns:
                df = df[['created_at', 'field5']].copy()
                df.columns = ['DateTime', 'Precipitation']
                df['DateTime'] = pd.to_datetime(df['DateTime'], utc=True).dt.tz_localize(None)  # 
                valid = clean_and_prepare(df, os.path.basename(file_path))
                dataframes.append(valid)
        except Exception as e:
            print(f"Error CSV {file_path}: {e}")

    if not dataframes:
        return None, None

    # --- Merge files ---
    consolidated_df = pd.concat(dataframes, ignore_index=True).sort_values('DateTime').reset_index(drop=True)

    # --- Fill gaps ---
    complete_df = fill_missing_timesteps(consolidated_df)

    complete_df = complete_df.rename(columns={'Precipitation': 'P_drops'})
    complete_df['P_mm'] = complete_df['P_drops'].apply(lambda x: x * 0.01 if x >= 0 else x)

    # --- Save CSV per station ---
    csv_file = os.path.join(output_directory, f"{station_name}_Consolidated_Complete.csv")
    complete_df.to_csv(csv_file, index=False)

    # --- Compute station summary statistics ---
    measurements = complete_df[complete_df['P_drops'] >= 0]
    precip_pos = measurements[measurements['P_mm'] > 0]['P_mm']

    nb_minus1 = (complete_df['P_drops'] == -1).sum()
    nb_minus2 = (complete_df['P_drops'] == -2).sum()
    total_rows = len(complete_df)
    nb_measurements = len(measurements)
    total_gaps = total_rows - nb_measurements
    pct_gaps = round(total_gaps / total_rows * 100, 1)                  # percentage of gaps in the complete dataset (including filled gaps)

    if len(precip_pos) > 0:
        mean_pos = round(precip_pos.mean(), 3)
        std_pos = round(precip_pos.std(), 3)
        median_pos = round(precip_pos.median(), 3)
        q95_pos = round(precip_pos.quantile(0.95), 3)
    else:
        mean_pos = std_pos = median_pos = q95_pos = 0

    summary_stats = {
        'Station': station_name,
        'Nb_rows_total': total_rows,
        'Nb_measurements': nb_measurements,
        'Nb_gaps_total': total_gaps,
        'Nb_-1_inter': nb_minus1,
        'Nb_-2_intra': nb_minus2,
        '%_gaps': pct_gaps,
        'Mean_pos': mean_pos,
        'Std_pos': std_pos,
        'Median_pos': median_pos,
        'Q95_pos': q95_pos
    }

    return complete_df, summary_stats


# %% 3 - Run all stations and collect summaries

all_summary = []

for folder_name in os.listdir(base_directory_processed):
    folder_path = os.path.join(base_directory_processed, folder_name)
    if os.path.isdir(folder_path):
        df_complete, summary_stats = process_consolidate(folder_path)
        if df_complete is not None:
            all_summary.append(summary_stats)

# --- Convert to DataFrame ---
df_summary = pd.DataFrame(all_summary)

# --- Sort by Station alphabetically ---
df_summary = df_summary.sort_values(by='Station').reset_index(drop=True)

# --- Save global summary ---
summary_file = os.path.join(output_directory, "SUMMARY_All_Stations.xlsx")
df_summary.to_excel(summary_file, index=False)
print("Global summary saved:", summary_file)


# %% 4 - Excel formatting

if os.path.exists(summary_file):
    wb = load_workbook(summary_file)
    ws = wb.active

    header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    fill_even = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    for row in range(2, ws.max_row + 1):
        fill = fill_even if row % 2 == 0 else fill_odd
        for col in range(1, ws.max_column + 1):
            ws.cell(row=row, column=col).fill = fill

    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2

    wb.save(summary_file)


# %%

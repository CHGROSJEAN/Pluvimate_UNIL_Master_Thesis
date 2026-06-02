# Charlotte Grosjean - 26.05.2026 - UNIL Master Thesis
"""
CROSS-CORRELATION ANALYSIS UTILITIES

Consolidated utility functions for cross-correlation analysis between MeteoSwiss and processed station data.

This module consolidates functions from:
- Archive/Stage_5_lausanne_cross_correlation.py

Main functions:
- parse_station_datetime(): Parse station timestamps with AM/PM fallback
- load_meteosuisse_source(): Load MeteoSwiss 10-minute precipitation data
- load_processed_file(): Load processed station files (TXT/XLSX/CSV)
- load_xlsx_precipitation(): Stream-read XLSX precipitation files
- infer_timestep_minutes(): Infer median time step from station data
- regularize_station_series(): Place station data on regular grid
- distribute_meteo_to_target_grid(): Redistribute 10-min MeteoSwiss to target grid
- pearson_corr_arrays(): Compute Pearson correlation with NaN filtering
- compute_best_cross_correlation(): Find lag with highest absolute correlation
"""

from __future__ import annotations

import os
import re
from datetime import datetime, time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from A_data_processing_utils import extract_numeric, parse_date


# ===== CONFIGURATION (imported from Stage_7) =====
# These constants are defined in their respective Stage files
# Functions import them locally to avoid circular dependencies


def parse_station_datetime(value):
    """
    Parse station timestamps, including TXT exports that use AM/PM with hour 0.

    data_processing_utils.parse_date already handles most local formats, but
    some files contain strings such as "Dec 12 2025 0:03:47 AM". Python's
    standard 12-hour parser rejects hour 0, so this fallback handles it
    explicitly.
    """
    parsed = parse_date(value)
    if pd.notna(parsed):
        return parsed

    if pd.isna(value):
        return pd.NaT

    text = str(value).strip().replace(",", "")
    match = re.match(
        r"^(?P<month>[A-Za-zéûû\.]+)\s+"
        r"(?P<day>\d{1,2})\s+"
        r"(?P<year>\d{4})\s+"
        r"(?P<hour>\d{1,2}):(?P<minute>\d{2}):(?P<second>\d{2})\s+"
        r"(?P<ampm>AM|PM)$",
        text,
        flags=re.IGNORECASE,
    )
    if not match:
        return pd.to_datetime(text, errors="coerce")

    month_map = {
        "jan": 1, "janv": 1,
        "feb": 2, "fev": 2, "févr": 2, "fevr": 2,
        "mar": 3, "mars": 3,
        "apr": 4, "avr": 4,
        "may": 5, "mai": 5,
        "jun": 6, "juin": 6,
        "jul": 7, "juil": 7,
        "aug": 8, "août": 8, "aout": 8,
        "sep": 9, "sept": 9,
        "oct": 10,
        "nov": 11,
        "dec": 12, "déc": 12,
    }

    month_key = match.group("month").replace(".", "").lower()
    month = month_map.get(month_key)
    if month is None:
        return pd.to_datetime(text, errors="coerce")

    hour = int(match.group("hour"))
    ampm = match.group("ampm").upper()
    if ampm == "AM":
        hour = 0 if hour in (0, 12) else hour
    else:
        hour = 12 if hour in (0, 12) else hour + 12

    try:
        return pd.Timestamp(
            year=int(match.group("year")),
            month=month,
            day=int(match.group("day")),
            hour=hour,
            minute=int(match.group("minute")),
            second=int(match.group("second")),
        )
    except ValueError:
        return pd.NaT


def load_meteosuisse_source(file_path: str) -> pd.DataFrame:
    """Load one MeteoSwiss 10-minute precipitation source."""
    df = pd.read_csv(file_path, sep=None, engine="python")

    required = {"reference_timestamp", "rre150z0"}
    missing = required.difference(df.columns)
    if missing:
        raise ValueError(f"MeteoSwiss file is missing columns: {sorted(missing)}")

    df = df[["reference_timestamp", "rre150z0"]].copy()
    df.columns = ["DateTime", "P_mm_10min"]
    df["DateTime"] = pd.to_datetime(df["DateTime"], dayfirst=True, errors="coerce")
    df["P_mm_10min"] = (
        df["P_mm_10min"]
        .astype(str)
        .str.replace(",", ".", regex=False)
        .pipe(pd.to_numeric, errors="coerce")
    )

    df = df.dropna(subset=["DateTime", "P_mm_10min"])
    df = df[df["P_mm_10min"] >= 0]
    df = df.sort_values("DateTime").drop_duplicates(subset=["DateTime"], keep="first")
    return df.reset_index(drop=True)


def load_meteosuisse_lausanne(file_path: str) -> pd.DataFrame:
    """Backward-compatible wrapper for older calls."""
    return load_meteosuisse_source(file_path)


def load_xlsx_precipitation(file_path: str) -> pd.DataFrame:
    """Read an XLSX processed file using openpyxl streaming mode."""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    worksheet = workbook.active

    rows = worksheet.iter_rows(min_row=7, values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        workbook.close()
        return pd.DataFrame(columns=["DateTime", "Precipitation"])

    columns = [str(value).strip() if value is not None else "" for value in header]
    lowered = [col.lower() for col in columns]

    date_idx = next(
        (idx for idx, col in enumerate(lowered) if "date" in col or "time" in col),
        0,
    )
    time_idx = next(
        (idx for idx, col in enumerate(lowered) if col in {"heure", "time"}),
        None,
    )
    precip_idx = next(
        (
            idx
            for idx, col in enumerate(lowered)
            if any(key in col for key in ["precip", "pulse", "pulsation"])
        ),
        len(columns) - 1,
    )

    datetimes = []
    precipitations = []
    for row in rows:
        date_value = row[date_idx] if date_idx < len(row) else None
        precip_value = row[precip_idx] if precip_idx < len(row) else None

        # Some exports split date and time. If the selected date column has no
        # time component, combine it with the explicit time column.
        if time_idx is not None and time_idx < len(row) and date_value is not None:
            time_value = row[time_idx]
            if isinstance(date_value, datetime) and isinstance(time_value, datetime):
                if date_value.time() == time(0, 0):
                    date_value = datetime.combine(date_value.date(), time_value.time())
            elif isinstance(date_value, datetime) and isinstance(time_value, time):
                if date_value.time() == time(0, 0):
                    date_value = datetime.combine(date_value.date(), time_value)

        datetimes.append(date_value)
        precipitations.append(precip_value)

    workbook.close()
    return pd.DataFrame({"DateTime": datetimes, "Precipitation": precipitations})


def load_processed_file(file_path: str) -> pd.DataFrame:
    """Load one processed station file as DateTime + P_mm."""
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".txt":
        df = pd.read_csv(
            file_path,
            sep="\t",
            skiprows=5,
            header=None,
            usecols=[0, 1, 2],
            names=["Index", "DateTime", "Precipitation"],
        )
        df["DateTime"] = df["DateTime"].apply(parse_station_datetime)

    elif extension == ".xlsx":
        if os.path.basename(file_path).startswith("~$"):
            return pd.DataFrame(columns=["DateTime", "P_mm"])

        df = load_xlsx_precipitation(file_path)
        df["DateTime"] = df["DateTime"].apply(parse_station_datetime)

    elif extension == ".csv":
        df = pd.read_csv(file_path)
        if "created_at" in df.columns and "field5" in df.columns:
            df = df[["created_at", "field5"]].copy()
            df.columns = ["DateTime", "Precipitation"]
            df["DateTime"] = pd.to_datetime(df["DateTime"], utc=True, errors="coerce").dt.tz_localize(None)
        elif {"DateTime", "P_mm"}.issubset(df.columns):
            df = df[["DateTime", "P_mm"]].copy()
            df["DateTime"] = pd.to_datetime(df["DateTime"], errors="coerce")
            df["Precipitation"] = df["P_mm"]
        else:
            raise ValueError("CSV format not recognized: expected created_at/field5 or DateTime/P_mm")

    else:
        raise ValueError(f"Unsupported file extension: {extension}")

    if "Precipitation" in df.columns:
        df["Precipitation"] = df["Precipitation"].apply(extract_numeric)

    df = df.dropna(subset=["DateTime", "Precipitation"]).copy()
    df = df[df["Precipitation"] >= 0]

    # Existing project convention: one pluviometer pulse/drop = 0.01 mm.
    if "P_mm" not in df.columns or extension != ".csv":
        df["P_mm"] = df["Precipitation"] * 0.01

    df = df[["DateTime", "P_mm"]]
    df = df.sort_values("DateTime")
    df = df.groupby("DateTime", as_index=False)["P_mm"].sum()
    return df.reset_index(drop=True)


def infer_timestep_minutes(series_df: pd.DataFrame) -> float | None:
    """Infer the median time step in minutes from one station file."""
    diffs = series_df["DateTime"].sort_values().diff().dropna()
    if diffs.empty:
        return None

    timestep = diffs.median().total_seconds() / 60
    if not np.isfinite(timestep) or timestep <= 0:
        return None

    rounded = round(timestep)
    if abs(timestep - rounded) < 0.05:
        return float(rounded)
    return float(timestep)


def regularize_station_series(series_df: pd.DataFrame, timestep_minutes: float) -> pd.Series:
    """
    Place the station file on a complete regular grid at its own time step.

    Some CSV timestamps drift by a few seconds over long periods. Instead of
    requiring exact equality with a perfect grid, each measurement is assigned
    to the nearest expected time step from the first timestamp.
    """
    df = series_df.sort_values("DateTime").set_index("DateTime")
    timestep = pd.Timedelta(minutes=timestep_minutes)
    start = df.index.min()

    # Divide timedeltas directly so pandas version changes in integer units
    # (ns/us) do not affect the grid assignment.
    grid_positions = np.rint((df.index - start) / timestep).astype(int)
    grid_times = start + pd.to_timedelta(grid_positions * timestep.value, unit="ns")

    gridded = (
        df.assign(GridTime=grid_times)
        .groupby("GridTime")["P_mm"]
        .sum()
        .sort_index()
    )
    target_index = pd.date_range(start=start, end=gridded.index.max(), freq=timestep)
    regularized = gridded.reindex(target_index).astype(float)
    if len(series_df) >= 10 and len(regularized) < len(series_df) * 0.5:
        raise ValueError(
            "Regularized station series is unexpectedly short "
            f"({len(regularized)} grid rows for {len(series_df)} source rows). "
            "This usually means an old timestamp-regularization function is being used. "
            "Restart the kernel and import the current H_cross_correlation_utils.py file."
        )
    return regularized


def distribute_meteo_to_target_grid(
    meteo_df: pd.DataFrame,
    target_index: pd.DatetimeIndex,
    target_step_minutes: float,
    meteoswiss_step_minutes: int = 10,
) -> pd.Series:
    """
    Distribute 10-minute MeteoSwiss totals onto a target grid.

    Both MeteoSwiss and target timestamps are treated as interval end times:
    source interval = (meteo_timestamp - 10 min, meteo_timestamp]
    target interval = (target_timestamp - target_step, target_timestamp]
    
    Args:
        meteo_df: MeteoSwiss data with DateTime and P_mm_10min columns
        target_index: Target time index
        target_step_minutes: Target time step in minutes
        meteoswiss_step_minutes: MeteoSwiss time step (default 10 minutes)
    """
    target_index = pd.DatetimeIndex(target_index)
    if len(target_index) == 0:
        return pd.Series(dtype=float)

    target_step = pd.Timedelta(minutes=target_step_minutes)
    source_step = pd.Timedelta(minutes=meteoswiss_step_minutes)
    source_step_ns = source_step.value

    target_ends = target_index
    target_starts = target_ends - target_step
    target_start_min = target_starts.min()
    target_end_max = target_ends.max()

    meteo = meteo_df.copy()
    meteo["interval_start"] = meteo["DateTime"] - source_step
    meteo["interval_end"] = meteo["DateTime"]
    meteo = meteo[
        (meteo["interval_end"] > target_start_min)
        & (meteo["interval_start"] < target_end_max)
    ]

    if meteo.empty:
        return pd.Series(np.nan, index=target_index, dtype=float)

    source_starts = meteo["interval_start"].astype("datetime64[ns]").astype("int64").to_numpy()
    source_ends = meteo["interval_end"].astype("datetime64[ns]").astype("int64").to_numpy()
    source_values = meteo["P_mm_10min"].to_numpy(dtype=float)

    target_starts_ns = target_starts.astype("datetime64[ns]").astype("int64").to_numpy()
    target_ends_ns = target_ends.astype("datetime64[ns]").astype("int64").to_numpy()

    values = np.zeros(len(target_index), dtype=float)
    has_data = np.zeros(len(target_index), dtype=bool)

    # Iterate over target bins. Each target bin overlaps only a few 10-minute
    # source intervals, so this is much faster than scanning all source rows.
    for pos, (target_start_ns, target_end_ns) in enumerate(zip(target_starts_ns, target_ends_ns)):
        first_source = np.searchsorted(source_ends, target_start_ns, side="right")
        last_source = np.searchsorted(source_starts, target_end_ns, side="left")

        for source_pos in range(first_source, last_source):
            overlap_start_ns = max(target_start_ns, source_starts[source_pos])
            overlap_end_ns = min(target_end_ns, source_ends[source_pos])
            overlap_ns = max(0, overlap_end_ns - overlap_start_ns)
            if overlap_ns > 0:
                values[pos] += source_values[source_pos] * (overlap_ns / source_step_ns)
                has_data[pos] = True

    values[~has_data] = np.nan
    return pd.Series(values, index=target_index, dtype=float)


def pearson_corr_arrays(
    station_values: np.ndarray,
    meteo_values: np.ndarray,
    min_valid_pairs: int = 10,
) -> tuple[float, int]:
    """
    Compute Pearson correlation for two same-length arrays with NaN filtering.
    
    Args:
        station_values: Station precipitation array
        meteo_values: MeteoSwiss precipitation array
        min_valid_pairs: Minimum required pairs for correlation calculation
    """
    mask = np.isfinite(station_values) & np.isfinite(meteo_values)
    n_pairs = int(mask.sum())
    if n_pairs < min_valid_pairs:
        return np.nan, n_pairs

    x = station_values[mask]
    y = meteo_values[mask]
    x_centered = x - x.mean()
    y_centered = y - y.mean()
    denominator = np.sqrt(np.sum(x_centered ** 2) * np.sum(y_centered ** 2))
    if denominator == 0:
        return np.nan, n_pairs

    return float(np.sum(x_centered * y_centered) / denominator), n_pairs


def compute_best_cross_correlation(
    station_series: pd.Series,
    meteo_series: pd.Series,
    timestep_minutes: float,
    max_lag_minutes: int = 240,
    min_valid_pairs: int = 10,
) -> dict:
    """
    Find the lag with the largest absolute Pearson correlation.
    
    Args:
        station_series: Station precipitation time series
        meteo_series: MeteoSwiss precipitation time series
        timestep_minutes: Time step in minutes for lag calculation
        max_lag_minutes: Maximum lag to test in minutes
        min_valid_pairs: Minimum valid pairs for correlation
    
    Returns:
        Dictionary with lag and correlation results
    """
    max_lag_steps = int(np.floor(max_lag_minutes / timestep_minutes))
    station_values = station_series.to_numpy(dtype=float)
    meteo_values = meteo_series.to_numpy(dtype=float)
    rows = []

    for lag_steps in range(-max_lag_steps, max_lag_steps + 1):
        if lag_steps > 0:
            # Same convention as meteo.shift(lag_steps):
            # station at t is compared with MeteoSwiss at t - lag.
            station_slice = station_values[lag_steps:]
            meteo_slice = meteo_values[:-lag_steps]
        elif lag_steps < 0:
            station_slice = station_values[:lag_steps]
            meteo_slice = meteo_values[-lag_steps:]
        else:
            station_slice = station_values
            meteo_slice = meteo_values

        corr, n_pairs = pearson_corr_arrays(station_slice, meteo_slice, min_valid_pairs)
        rows.append(
            {
                "lag_steps": lag_steps,
                "lag_minutes": lag_steps * timestep_minutes,
                "correlation": corr,
                "n_valid_pairs": n_pairs,
            }
        )

    lag_df = pd.DataFrame(rows).dropna(subset=["correlation"])
    if lag_df.empty:
        return {
            "best_lag_steps": np.nan,
            "best_lag_minutes": np.nan,
            "best_correlation": np.nan,
            "best_abs_correlation": np.nan,
            "n_valid_pairs_at_best_lag": 0,
            "_lag_df": lag_df,
        }

    best = lag_df.loc[lag_df["correlation"].abs().idxmax()]
    return {
        "best_lag_steps": int(best["lag_steps"]),
        "best_lag_minutes": float(best["lag_minutes"]),
        "best_correlation": float(best["correlation"]),
        "best_abs_correlation": float(abs(best["correlation"])),
        "n_valid_pairs_at_best_lag": int(best["n_valid_pairs"]),
        "_lag_df": lag_df,
    }
